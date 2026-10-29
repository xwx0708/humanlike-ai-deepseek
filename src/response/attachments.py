"""读用户发来的文件和图片（D38）。用户原话："增加分析功能让它可以读文件、读图片等等。支持不包含 MP4 在内的大多数文件。"

- 图片：交给思考模型直接看（模型要支持看图）。
- 文件：在这里把内容提取成文字，放进这一轮的提示词；纯文本、代码、PDF、Word、Excel、PPT、压缩包（只列文件名）等。
- 不支持：视频（包括 MP4）、音频、旧版 Office 二进制格式（.doc / .xls / .ppt，请另存为新格式）、其他二进制文件。
  不支持时告诉模型，让它如实跟用户说。
- 记进记忆的不是全文：只记"发来了什么"和开头一段（暂定 2000 字），免得一个大文件挤满记忆。
数值都是暂定的。
"""

import zipfile
from dataclasses import dataclass
from pathlib import Path

from runtime import config

MAX_BYTES = int(config.value("attachments", "max_mb", 20.0) * 1024 * 1024)  # 单个文件上限（暂定）
PROMPT_CHARS = config.value("attachments", "prompt_chars", 20000)  # 一个文件最多放进提示词多少字（暂定）
MEMORY_CHARS = config.value("attachments", "memory_chars", 2000)  # 记进记忆的开头部分（暂定）

IMAGE = (".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp")
TEXT = (
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg",
    ".conf", ".log", ".xml", ".html", ".htm", ".css", ".srt", ".ass", ".tex", ".rst", ".sql", ".sh", ".bat",
    ".ps1", ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".kt", ".c", ".h", ".cpp", ".hpp", ".cs", ".go",
    ".rs", ".rb", ".php", ".swift", ".lua", ".r", ".m", ".vue", ".dart", ".scala",
)  # fmt: skip
VIDEO = (".mp4", ".mov", ".avi", ".mkv", ".flv", ".wmv", ".webm", ".m4v", ".3gp")
AUDIO = (".mp3", ".wav", ".flac", ".aac", ".ogg", ".m4a", ".amr", ".silk", ".wma")
OLD_OFFICE = (".doc", ".xls", ".ppt")


@dataclass
class Attachment:
    name: str
    kind: str  # "图片" / "文件" / "不支持"
    text: str = ""  # 提取出的文字（文件）
    image_path: str = ""  # 本地路径（图片）
    reason: str = ""  # 不支持或读不出的原因
    truncated: bool = False

    def prompt_block(self) -> str:
        """放进这一轮提示词的部分。"""
        if self.kind == "图片":
            return f"（对方发来一张图片：{self.name}，图片已附上）"
        if self.kind == "不支持":
            return f"（对方发来文件「{self.name}」，但你打不开：{self.reason}。如实告诉对方。）"
        more = f"\n（只看了前 {PROMPT_CHARS} 字，后面还有）" if self.truncated else ""
        return f"（对方发来文件「{self.name}」，内容如下）\n{self.text}{more}"

    def memory_note(self) -> str:
        """记进记忆的部分：发来了什么 + 开头一段。"""
        if self.kind == "图片":
            return f"（发来一张图片：{self.name}）"
        if self.kind == "不支持":
            return f"（发来文件「{self.name}」，打不开：{self.reason}）"
        head = self.text[:MEMORY_CHARS]
        return f"（发来文件「{self.name}」，开头是：{head}{'……' if len(self.text) > MEMORY_CHARS else ''}）"


def read_image(path: str | Path, name: str = "") -> Attachment:
    return Attachment(name or Path(path).name, "图片", image_path=str(path))


def read_file(path: str | Path, name: str = "") -> Attachment:
    p = Path(path)
    name = name or p.name
    suffix = Path(name).suffix.lower() or p.suffix.lower()
    if suffix in IMAGE:
        return read_image(p, name)
    if suffix in VIDEO:
        return unsupported(name, "视频还看不了")
    if suffix in AUDIO:
        return unsupported(name, "音频还听不了")
    if suffix in OLD_OFFICE:
        return unsupported(name, f"旧版 Office 格式（{suffix}）读不了，另存为 {suffix}x 再发")
    if not p.is_file():
        return unsupported(name, "文件没下载下来")
    if p.stat().st_size > MAX_BYTES:
        return unsupported(name, f"文件太大（超过 {MAX_BYTES // 1024 // 1024} MB）")
    try:
        if suffix == ".pdf":
            text = _pdf(p)
        elif suffix in (".docx", ".docm"):
            text = _docx(p)
        elif suffix in (".xlsx", ".xlsm"):
            text = _xlsx(p)
        elif suffix == ".pptx":
            text = _pptx(p)
        elif suffix == ".zip":
            text = _zip(p)
        elif suffix in TEXT or _looks_like_text(p):
            text = _decode(p.read_bytes())
        else:
            return unsupported(name, f"不认识的文件类型（{suffix or '没有扩展名'}）")
    except Exception as e:  # noqa: BLE001  文件坏了、加密了等，都如实告诉模型
        return unsupported(name, f"读的时候出错了（{type(e).__name__}）")
    text = text.strip()
    if not text:
        return unsupported(name, "里面没有能读出的文字（可能是扫描件或纯图片）")
    return Attachment(name, "文件", text=text[:PROMPT_CHARS], truncated=len(text) > PROMPT_CHARS)


def unsupported(name: str, reason: str) -> Attachment:
    """打不开的附件：告诉模型原因，让它如实跟对方说。"""
    return Attachment(name, "不支持", reason=reason)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _looks_like_text(p: Path) -> bool:
    head = p.read_bytes()[:4096]
    return b"\x00" not in head and bool(head)


def _pdf(p: Path) -> str:
    from pypdf import PdfReader

    pages = PdfReader(str(p)).pages
    return "\n\n".join(f"[第 {i} 页]\n{page.extract_text() or ''}" for i, page in enumerate(pages, 1))


def _docx(p: Path) -> str:
    import docx

    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs]
    for table in d.tables:
        parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    return "\n".join(parts)


def _xlsx(p: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(p), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[工作表 {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                parts.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def _pptx(p: Path) -> str:
    from pptx import Presentation

    parts = []
    for i, slide in enumerate(Presentation(str(p)).slides, 1):
        parts.append(f"[第 {i} 页]")
        parts += [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    return "\n".join(parts)


def _zip(p: Path) -> str:
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
    return f"压缩包里有 {len(names)} 个文件（只列文件名，没有解压读内容）：\n" + "\n".join(names[:200])
