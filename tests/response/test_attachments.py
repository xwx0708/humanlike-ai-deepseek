"""读用户发来的文件和图片（D38）：大多数文件能读，MP4 等视频、音频不支持。"""

import zipfile

import pytest

from memory.scanner import Scanner
from memory.store import MemoryStore
from response import attachments as att
from response.attachments import read_file
from response.responder import Responder
from runtime.lifecycle import RUNNING


def minimal_pdf(text: str) -> bytes:
    """按 PDF 规范拼一个只有一页、一行英文的文件（带交叉引用表）。"""
    stream = f"BT /F1 18 Tf 10 50 Td ({text}) Tj ET".encode()
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 100]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n" % len(stream) + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj" % i + body + b"endobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


def test_plain_text_and_code_including_gbk(tmp_path):
    (tmp_path / "笔记.txt").write_bytes("今天学了遗忘曲线".encode("gb18030"))
    (tmp_path / "a.py").write_text("print('hi')", encoding="utf-8")
    assert read_file(tmp_path / "笔记.txt").text == "今天学了遗忘曲线"
    assert read_file(tmp_path / "a.py").kind == "文件"


def test_office_files(tmp_path):
    import docx
    from openpyxl import Workbook
    from pptx import Presentation

    d = docx.Document()
    d.add_paragraph("周报：完成了记忆模块")
    d.save(tmp_path / "周报.docx")
    wb = Workbook()
    wb.active.append(["名字", "分数"])
    wb.active.append(["小d", 99])
    wb.save(tmp_path / "表.xlsx")
    p = Presentation()
    p.slides.add_slide(p.slide_layouts[1]).shapes.title.text = "第一页标题"
    p.save(tmp_path / "汇报.pptx")
    assert "完成了记忆模块" in read_file(tmp_path / "周报.docx").text
    assert "小d | 99" in read_file(tmp_path / "表.xlsx").text
    assert "第一页标题" in read_file(tmp_path / "汇报.pptx").text


def test_pdf_and_zip(tmp_path):
    (tmp_path / "x.pdf").write_bytes(minimal_pdf("Hello PDF"))
    assert "Hello PDF" in read_file(tmp_path / "x.pdf").text
    with zipfile.ZipFile(tmp_path / "包.zip", "w") as z:
        z.writestr("里面/readme.md", "hi")
    assert "里面/readme.md" in read_file(tmp_path / "包.zip").text


def test_unsupported_types_say_why(tmp_path):
    for name in ("片段.mp4", "歌.mp3", "旧.doc"):
        (tmp_path / name).write_bytes(b"\x00\x01")
    assert "视频" in read_file(tmp_path / "片段.mp4").reason
    assert "音频" in read_file(tmp_path / "歌.mp3").reason
    assert ".docx" in read_file(tmp_path / "旧.doc").reason
    (tmp_path / "blob.bin").write_bytes(b"\x00\xff" * 10)
    assert read_file(tmp_path / "blob.bin").kind == "不支持"
    assert "没下载下来" in read_file(tmp_path / "不存在.txt").reason


def test_images_are_passed_through(tmp_path):
    (tmp_path / "猫.png").write_bytes(b"png")
    a = read_file(tmp_path / "猫.png")
    assert a.kind == "图片" and a.image_path.endswith("猫.png")


def test_long_text_truncated_for_prompt_and_memory(tmp_path, monkeypatch):
    monkeypatch.setattr(att, "PROMPT_CHARS", 50)
    monkeypatch.setattr(att, "MEMORY_CHARS", 10)
    (tmp_path / "长.txt").write_text("字" * 80, encoding="utf-8")
    a = read_file(tmp_path / "长.txt")
    assert a.truncated and len(a.text) == 50 and "后面还有" in a.prompt_block()
    assert a.memory_note() == "（发来文件「长.txt」，开头是：" + "字" * 10 + "……）"


@pytest.fixture
def store():
    s = MemoryStore(":memory:")
    yield s
    s.close()


def test_responder_puts_full_text_in_prompt_but_only_a_note_in_memory(store, tmp_path, monkeypatch):
    monkeypatch.setattr(att, "MEMORY_CHARS", 4)
    (tmp_path / "需求.txt").write_text("要做一个会遗忘的 AI", encoding="utf-8")
    (tmp_path / "图.png").write_bytes(b"png")
    r = Responder(store, Scanner(store), lambda: "", phase=lambda: RUNNING)
    p = r.on_user_message("帮我看看", 0, [read_file(tmp_path / "需求.txt"), read_file(tmp_path / "图.png")])
    assert "要做一个会遗忘的 AI" in p.prompt and "图片已附上" in p.prompt
    assert p.image_paths == [str(tmp_path / "图.png")]
    saved = store.all_turns()[-1].content
    assert saved.startswith("帮我看看\n（发来文件「需求.txt」，开头是：要做一个……）")
    assert "（发来一张图片：图.png）" in saved
